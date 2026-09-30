import os
import time
from typing import Dict, Any
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from utils import format_seconds

class PPTXGenerator:
    def __init__(self, output_dir: str = "exports"):
        self.output_dir = output_dir
        os.makedirs(self.output_dir, exist_ok=True)

        # Color Palette
        self.COLOR_NAVY = RGBColor(15, 23, 42)       # #0F172A
        self.COLOR_INDIGO = RGBColor(79, 70, 229)    # #4F46E5
        self.COLOR_BG_LIGHT = RGBColor(248, 250, 252)# #F8FAFC
        self.COLOR_TEXT_DARK = RGBColor(30, 41, 59)  # #1E293B
        self.COLOR_MUTED = RGBColor(100, 116, 139)   # #64748B
        self.COLOR_WHITE = RGBColor(255, 255, 255)
        self.COLOR_CARD_BG = RGBColor(255, 255, 255)

    def generate_presentation(self, recording: Dict[str, Any], output_filename: str = None) -> str:
        """Generates a professional 16:9 PowerPoint presentation from session summary."""
        title = recording.get("title", "Training Program")
        rec_id = recording.get("id", 1)
        duration = recording.get("duration", 0)
        recorded_at = recording.get("recorded_at", time.strftime("%Y-%m-%d"))
        summary = recording.get("summary", {})

        if not output_filename:
            safe_title = "".join(c for c in title if c.isalnum() or c in (' ', '_', '-')).rstrip()
            safe_title = safe_title.replace(' ', '_')
            output_filename = f"TrainIQ_Presentation_{rec_id}_{safe_title}.pptx"

        filepath = os.path.join(self.output_dir, output_filename)

        prs = Presentation()
        prs.slide_width = Inches(13.333)
        prs.slide_height = Inches(7.5)
        blank_layout = prs.slide_layouts[6]  # Blank slide

        # -----------------------------------------------------------------
        # SLIDE 1: Title Slide (Dark Navy Background)
        # -----------------------------------------------------------------
        slide1 = prs.slides.add_slide(blank_layout)
        self._add_full_background(slide1, self.COLOR_NAVY)

        # Accent Bar
        accent_bar = slide1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1.0), Inches(1.5), Inches(0.15), Inches(4.5))
        accent_bar.fill.solid()
        accent_bar.fill.fore_color.rgb = self.COLOR_INDIGO
        accent_bar.line.fill.background()

        # Title & Subtitle Box
        title_box = slide1.shapes.add_textbox(Inches(1.5), Inches(1.8), Inches(10.5), Inches(3.8))
        tf = title_box.text_frame
        tf.word_wrap = True
        
        p0 = tf.paragraphs[0]
        p0.text = "TRAINIQ EXECUTIVE PRESENTATION"
        p0.font.size = Pt(14)
        p0.font.bold = True
        p0.font.color.rgb = self.COLOR_INDIGO
        p0.space_after = Pt(12)

        p1 = tf.add_paragraph()
        p1.text = title
        p1.font.size = Pt(36)
        p1.font.bold = True
        p1.font.color.rgb = self.COLOR_WHITE
        p1.space_after = Pt(20)

        p2 = tf.add_paragraph()
        p2.text = f"Recorded Date: {recorded_at}   |   Session Duration: {format_seconds(duration)}   |   ID: #{rec_id}"
        p2.font.size = Pt(14)
        p2.font.color.rgb = self.COLOR_MUTED

        # -----------------------------------------------------------------
        # SLIDE 2: Executive Summary
        # -----------------------------------------------------------------
        slide2 = prs.slides.add_slide(blank_layout)
        self._add_full_background(slide2, self.COLOR_BG_LIGHT)
        self._add_slide_header(slide2, "Executive Overview", "High-level summary of key discussion points and objectives")

        # Summary Card
        card = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.0), Inches(1.8), Inches(11.333), Inches(4.8))
        card.fill.solid()
        card.fill.fore_color.rgb = self.COLOR_CARD_BG
        card.line.color.rgb = RGBColor(226, 232, 240)

        tf2 = card.text_frame
        tf2.word_wrap = True
        tf2.margin_left = Inches(0.4)
        tf2.margin_right = Inches(0.4)
        tf2.margin_top = Inches(0.4)

        p_sum_head = tf2.paragraphs[0]
        p_sum_head.text = "Session Summary & Core Intent"
        p_sum_head.font.size = Pt(20)
        p_sum_head.font.bold = True
        p_sum_head.font.color.rgb = self.COLOR_INDIGO
        p_sum_head.space_after = Pt(14)

        p_sum_body = tf2.add_paragraph()
        p_sum_body.text = summary.get("executive_summary", "No summary available.")
        p_sum_body.font.size = Pt(16)
        p_sum_body.font.color.rgb = self.COLOR_TEXT_DARK

        # -----------------------------------------------------------------
        # SLIDE 3: Key Agenda Topics & Timestamp Breakdown
        # -----------------------------------------------------------------
        topics = summary.get("topics", [])
        if topics:
            slide3 = prs.slides.add_slide(blank_layout)
            self._add_full_background(slide3, self.COLOR_BG_LIGHT)
            self._add_slide_header(slide3, "Agenda & Key Topics", "Timestamped section breakdown recorded during training")

            top_lefts = [
                (Inches(1.0), Inches(1.8)),
                (Inches(6.8), Inches(1.8)),
                (Inches(1.0), Inches(4.3)),
                (Inches(6.8), Inches(4.3))
            ]

            for idx, top in enumerate(topics[:4]):
                pos_x, pos_y = top_lefts[idx]
                t_card = slide3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, pos_x, pos_y, Inches(5.5), Inches(2.2))
                t_card.fill.solid()
                t_card.fill.fore_color.rgb = self.COLOR_CARD_BG
                t_card.line.color.rgb = RGBColor(226, 232, 240)

                tf_t = t_card.text_frame
                tf_t.word_wrap = True
                tf_t.margin_left = Inches(0.25)
                tf_t.margin_top = Inches(0.2)

                p_ts = tf_t.paragraphs[0]
                p_ts.text = f"{top.get('timestamp', '[00:00]')} - {top.get('topic', 'Topic')}"
                p_ts.font.size = Pt(15)
                p_ts.font.bold = True
                p_ts.font.color.rgb = self.COLOR_INDIGO
                p_ts.space_after = Pt(6)

                p_det = tf_t.add_paragraph()
                p_det.text = top.get("details", "")
                p_det.font.size = Pt(12)
                p_det.font.color.rgb = self.COLOR_TEXT_DARK

        # -----------------------------------------------------------------
        # SLIDE 4: Key Learning Takeaways
        # -----------------------------------------------------------------
        takeaways = summary.get("key_takeaways", [])
        if takeaways:
            slide4 = prs.slides.add_slide(blank_layout)
            self._add_full_background(slide4, self.COLOR_BG_LIGHT)
            self._add_slide_header(slide4, "Key Learning Takeaways", "Essential conclusions and insights from the program")

            tk_box = slide4.shapes.add_textbox(Inches(1.0), Inches(1.8), Inches(11.333), Inches(5.0))
            tf_tk = tk_box.text_frame
            tf_tk.word_wrap = True

            for idx, tk in enumerate(takeaways[:5]):
                p = tf_tk.add_paragraph() if idx > 0 else tf_tk.paragraphs[0]
                p.text = f"•   {tk}"
                p.font.size = Pt(16)
                p.font.color.rgb = self.COLOR_TEXT_DARK
                p.space_after = Pt(16)

        # -----------------------------------------------------------------
        # SLIDE 5: Action Items & Follow-Up Tasks
        # -----------------------------------------------------------------
        actions = summary.get("action_items", [])
        if actions:
            slide5 = prs.slides.add_slide(blank_layout)
            self._add_full_background(slide5, self.COLOR_BG_LIGHT)
            self._add_slide_header(slide5, "Action Items & Assignments", "Tasks and follow-up activities assigned during training")

            act_card = slide5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.0), Inches(1.8), Inches(11.333), Inches(4.8))
            act_card.fill.solid()
            act_card.fill.fore_color.rgb = self.COLOR_CARD_BG
            act_card.line.color.rgb = RGBColor(226, 232, 240)

            tf_act = act_card.text_frame
            tf_act.word_wrap = True
            tf_act.margin_left = Inches(0.4)
            tf_act.margin_top = Inches(0.3)

            for idx, act in enumerate(actions[:5]):
                p = tf_act.add_paragraph() if idx > 0 else tf_act.paragraphs[0]
                p.text = f"[  ]   {act}"
                p.font.size = Pt(16)
                p.font.bold = True
                p.font.color.rgb = self.COLOR_INDIGO
                p.space_after = Pt(16)

        prs.save(filepath)
        return filepath

    def _add_full_background(self, slide, color: RGBColor):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
        bg.fill.solid()
        bg.fill.fore_color.rgb = color
        bg.line.fill.background()

    def _add_slide_header(self, slide, title: str, subtitle: str):
        header_box = slide.shapes.add_textbox(Inches(1.0), Inches(0.5), Inches(11.333), Inches(1.2))
        tf = header_box.text_frame
        tf.word_wrap = True
        
        p0 = tf.paragraphs[0]
        p0.text = title
        p0.font.size = Pt(24)
        p0.font.bold = True
        p0.font.color.rgb = self.COLOR_NAVY
        p0.space_after = Pt(4)

        p1 = tf.add_paragraph()
        p1.text = subtitle
        p1.font.size = Pt(12)
        p1.font.color.rgb = self.COLOR_MUTED
